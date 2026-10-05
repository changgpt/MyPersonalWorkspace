import io

import docx

from daybook.docx_utils import docx_bytes_to_markdown


def _build_docx():
    document = docx.Document()
    document.add_heading("Meeting notes", level=1)
    document.add_paragraph("Plain paragraph text.")
    bullet = document.add_paragraph(style="List Bullet")
    bullet.add_run("First item")
    document.add_heading("Decisions", level=2)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_headings_converted_to_markdown_hashes():
    markdown_text = docx_bytes_to_markdown(_build_docx())
    assert "# Meeting notes" in markdown_text
    assert "## Decisions" in markdown_text


def test_plain_paragraph_preserved():
    markdown_text = docx_bytes_to_markdown(_build_docx())
    assert "Plain paragraph text." in markdown_text


def test_bullet_list_item_converted():
    markdown_text = docx_bytes_to_markdown(_build_docx())
    assert "- First item" in markdown_text
